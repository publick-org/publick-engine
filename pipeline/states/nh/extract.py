"""Save New Hampshire's statewide figures from the state's files, once a year.

New Hampshire's Department of Revenue Administration (DRA) and Department of
Education publish each year's tax rates, school figures, and cost per pupil as
statewide spreadsheets, one row per town or district. Their websites refuse
automated requests, so the files are downloaded by hand in a browser (about
once a year, when a new year is posted) and this command reads them. It keeps
only the figures the sites show, for every town and district, in
pipeline/states/nh/figures/, and each New Hampshire town's daily run reads its
own rows from there. Years already saved are kept; a year in a new file
replaces the same year saved before.

Download from:

  DRA, Municipal and Village District Tax Rates (DRA_PAGE below):
    <year> municipal and village (district) tax rates (.xlsx)
    <year> tax rate calculation data (.xlsx)
  Department of Education, Assessment Data (ASSESSMENT_PAGE):
    Public disaggregated data, CSV format (one per spring)
  Department of Education, Financial Reports (FINANCE_PAGE):
    Cost per pupil by district (.csv), one per fiscal year
  Department of Education iPlatform (GRADUATION_PAGE), Performance Data >
  Dropouts and Completers > Cohort Counts By School, exported as CSV for each
  school year: the export's address is GRADUATION_EXPORT with SchoolYear set.

Then, from the engine's folder:

    python -m pipeline.states.nh.extract ~/Downloads/*.xlsx ~/Downloads/*.csv
    python -m pipeline.states.nh.extract --population    # Census estimates, fetched directly

Each file is recognized by what's in it, not its name (except a tax rate
calculation file, whose year is only in its name). Commit the changed files in
figures/; the next engine release takes them to every New Hampshire town.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

FIGURES_DIR = Path(__file__).resolve().parent / "figures"

DRA_PAGE = ("https://www.revenue.nh.gov/about-dra/municipal-and-property-division/"
            "municipal-and-property-reports/municipal-and-village")
EDUCATION_STATISTICS = ("https://www.education.nh.gov/who-we-are/division-of-educator-and-analytic-resources/"
                        "bureau-of-education-statistics")
ASSESSMENT_PAGE = f"{EDUCATION_STATISTICS}/assessment-data"
FINANCE_PAGE = f"{EDUCATION_STATISTICS}/financial-reports"
GRADUATION_PAGE = "https://www.education.nh.gov/who-we-are/division-of-educator-and-analytic-resources/iplatform"
GRADUATION_EXPORT = ("https://my.doe.nh.gov/iPlatform/Report/ExportReport/?reportPath=/BDMQ/iPlatform Reports/"
                     "Performance Data/Dropouts and Completers/Cohort Counts By School&format=CSV&ReportID=30"
                     "&ReportViewerEnablePaging=True&SchoolYear=<year>")
POPULATION_PAGE = "https://www.census.gov/data/tables/time-series/demo/popest/2020s-total-cities-and-towns.html"
POPULATION_URL = ("https://www2.census.gov/programs-surveys/popest/datasets/2020-{year}/cities/totals/"
                  "sub-est{year}_33.csv")

# The state's own name for its totals, used as the key next to each town's or district's.
STATE = "State"

# Each figures file: what it holds, where it's from.
FILES = {
    "tax": {"source": "New Hampshire Department of Revenue Administration, Municipal and Property Division",
            "source_url": DRA_PAGE},
    "graduation": {"source": "New Hampshire Department of Education, cohort graduation rates (iPlatform)",
                   "source_url": GRADUATION_PAGE},
    "assessment": {"source": "New Hampshire Department of Education, New Hampshire Statewide Assessment System",
                   "source_url": ASSESSMENT_PAGE},
    "cost_per_pupil": {"source": "New Hampshire Department of Education, Bureau of School Finance",
                       "source_url": FINANCE_PAGE},
    "population": {"source": "U.S. Census Bureau, Population Estimates Program (city and town totals)",
                   "source_url": POPULATION_PAGE},
}


def norm(text) -> str:
    """A header as plain lowercase words: 'Valuation\\nIncuding\\nUtilities' -> 'valuation including utilities'."""
    return re.sub(r"\s+", " ", str(text or "")).strip().lower().replace("incuding", "including")


def number(value) -> float | None:
    """123, '12,980', '$23,533.31 ', '87.54%', ' -  ' -> a number, or None for a blank or dash."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("$", "").replace(",", "").replace("%", "").strip()
    if text in ("", "-", "–"):
        return None
    return float(text)


def whole(value) -> int | None:
    n = number(value)
    return None if n is None else round(n)


# ---- DRA: tax rates and the tax rate calculation ----

RATE_COLUMNS = {
    "date": "set_on", "valuation": "valuation", "valuation including utilities": "valuation_with_utilities",
    "municipal tax rate": "municipal", "county tax rate": "county", "state education tax rate": "state_education",
    "local education tax rate": "local_education", "total tax rate": "total", "total commitment": "commitment",
}
CALCULATION_COLUMNS = {
    "town appropriation": "town_appropriation", "town tax effort": "town_tax_effort",
    "net local school appropriation": "local_school_appropriation",
    "net local school appropriations": "local_school_appropriation",
    "net cooperative school #1 apportionment": "cooperative_school_apportionment",
    "net cooperative school #2 apportionment": "cooperative_school_apportionment",
    "local school tax effort": "local_school_tax_effort", "state education tax effort": "state_education_tax_effort",
    "county tax effort": "county_tax_effort", "total tax effort": "total_tax_effort",
    "(minus) veterans' credits": "veterans_credits", "(minus) veterans' tax credits": "veterans_credits",
    "plus village district(s) tax effort": "village_tax_effort", "(plus) village district(s) tax effort": "village_tax_effort",
    "total tax commitment": "commitment",
}
RATES = ("municipal", "county", "state_education", "local_education", "total")


def dra_table(content: bytes, columns: dict) -> dict[str, dict]:
    """{municipality: {field: value}} from a DRA workbook's first sheet.

    The header row is the one naming the municipality column. Columns move between
    years (merged cells, blank spacer columns), so each is found by its header.
    """
    sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows)
                 if any(norm(c).startswith("municipality") or "unincorporated place" in norm(c) for c in r))
    header = {j: norm(c) for j, c in enumerate(rows[start]) if c not in (None, "")}
    name_col = min(header)
    fields = {j: columns[h] for j, h in header.items() if h in columns}
    out = {}
    for row in rows[start + 1:]:
        name = row[name_col] if name_col < len(row) else None
        # The calculation files from 2023 on have a row of placeholder names under the header.
        if not isinstance(name, str) or not name.strip() or name.strip().lower().startswith("column"):
            continue
        record: dict = {}
        for j, field in fields.items():
            value = row[j] if j < len(row) else None
            if field == "set_on":
                record[field] = value.date().isoformat() if isinstance(value, datetime) else None
            elif field in RATES:
                record[field] = number(value)
            else:
                n = whole(value)
                if record.get(field) is not None:
                    # Both cooperative school columns add up to one figure.
                    record[field] += n or 0
                else:
                    record[field] = n
        out[re.sub(r"\s+", " ", name).strip()] = record
    return out


def dra_year(content: bytes, name: str) -> int:
    """The tax year: from the sheet's name ('2025 Municipal Tax Rates'), else the file's name."""
    sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True).sheetnames[0]
    found = re.match(r"(20\d\d)\b", sheet) or re.match(r"(20\d\d)\b", Path(name).name)
    if not found:
        raise SystemExit(f"{name}: can't tell the tax year. Name the file as the DRA does, starting with the year.")
    return int(found.group(1))


# ---- Department of Education ----

def graduation(text: str) -> tuple[int, dict]:
    """(class year, {district: {cohort, graduated, rate}, "State": {...}}) from an iPlatform cohort export."""
    title = re.search(r"(\d{4}) - (\d{4}) Cohort Graduation", text)
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("Textbox15,"))
    rows = list(csv.DictReader(io.StringIO("\n".join(lines[start:]))))
    out = {}
    for r in rows:
        if r.get("DstName"):
            out[r["DstName"].strip()] = {"cohort": whole(r["adjustedcohort3"]), "graduated": whole(r["graduated2"]),
                                         "rate": number(r["gradrate3"])}
    if rows:
        r = rows[0]
        out[STATE] = {"cohort": whole(r["adjustedcohort"]), "graduated": whole(r["graduated"]), "rate": number(r["Textbox8"])}
    return int(title.group(2)), out


def percent(value: str) -> int | None:
    """A published share: '31' -> 31. Suppressed or banded values ('< 10 %', '*') -> None."""
    text = (value or "").strip()
    return int(text) if text.isdigit() else None


def assessment(text: str) -> dict[int, dict]:
    """{spring: {district: {subject: {proficient, participation, students}}}} for all grades and all students.

    proficient is the published share scoring at levels 3 and 4 ("Above prof%").
    """
    out: dict[int, dict] = {}
    for r in csv.DictReader(io.StringIO(text)):
        level = r["Level of Data"].strip()
        if level not in ("District Level", "State Level") or r["Grade"].strip().lower() != "all grades":
            continue
        if r["Subgroup"].strip() != "All students" or r.get("DenominatorType", "Regular").strip() != "Regular":
            continue
        # The 2019 file has rows whose year a spreadsheet turned into a date ("11-Jul").
        if not r["yearid"].strip().isdigit():
            continue
        name = STATE if level == "State Level" else r["District"].strip()
        out.setdefault(int(r["yearid"]), {}).setdefault(name, {})[r["Subject"].strip()] = {
            "proficient": percent(r["Above prof% (lvl 3&4)"]), "participation": percent(r["Participate%"]),
            "students": r["Total FAY Students"].strip()}
    return out


def cost_per_pupil(text: str) -> tuple[int, dict]:
    """(fiscal year, {district: total cost per pupil, "State": the state average}) from the cost per pupil report."""
    title = re.search(r"COST PER PUPIL BY DISTRICT,\s*(\d{4})-(\d{4})", text, re.I)
    rows = list(csv.reader(io.StringIO(text)))
    head = next(i for i, r in enumerate(rows) if len(r) > 3 and r[0].strip() == "DIST")
    out = {}
    for r in rows[head + 1:]:
        if len(r) < 8 or not r[3].strip():
            continue
        name = r[3].strip()
        if name.lower().startswith("state average"):
            out[STATE] = number(r[7])
        elif r[0].strip().isdigit():
            out[name] = number(r[7])
    return int(title.group(2)), out


# ---- Census ----

def population(text: str) -> dict[int, dict]:
    """{year: {town: estimate}} for New Hampshire's towns and cities ('Manchester city' -> 'Manchester')."""
    out: dict[int, dict] = {}
    for r in csv.DictReader(io.StringIO(text)):
        if r["SUMLEV"] != "061" or not r["NAME"].endswith((" town", " city")):
            continue
        name = r["NAME"].rsplit(" ", 1)[0]
        for key, value in r.items():
            if key.startswith("POPESTIMATE") and key[11:].isdigit():
                out.setdefault(int(key[11:]), {})[name] = int(value)
    return out


# ---- Saving ----

def load(kind: str) -> dict:
    path = FIGURES_DIR / f"{kind}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {**FILES[kind], "years": {}}


def save(kind: str, data: dict, now: datetime) -> None:
    data.update(FILES[kind], extracted_at=now.isoformat(timespec="seconds"))
    data["years"] = dict(sorted(data["years"].items()))
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    (FIGURES_DIR / f"{kind}.json").write_text(dump(data), encoding="utf-8")


def dump(data: dict) -> str:
    """JSON with one line per town or district, so a new year's file changes little else and stays small."""
    compact = lambda v: json.dumps(v, ensure_ascii=False, separators=(",", ":"))
    head = [f" {compact(k)}: {compact(v)}" for k, v in data.items() if k != "years"]
    years = [f"  {compact(year)}: {{\n" + ",\n".join(f"   {compact(name)}: {compact(row)}" for name, row in rows.items())
             + "\n  }" for year, rows in data["years"].items()]
    return "{\n" + ",\n".join(head + [' "years": {\n' + ",\n".join(years) + "\n }"]) + "\n}\n"


def read(path: Path) -> tuple[str, dict[int, dict]]:
    """(figures file, {year: {name: record}}) for one downloaded file, recognized by its contents."""
    content = path.read_bytes()
    if content.startswith(b"PK"):
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        cells = " ".join(norm(c) for row in sheet.iter_rows(max_row=8, values_only=True) for c in row if c)
        year = dra_year(content, path.name)
        if "town appropriation" in cells:
            return "tax", {year: dra_table(content, CALCULATION_COLUMNS)}
        if "total tax rate" in cells or "municipal tax rate" in cells:
            return "tax", {year: dra_table(content, RATE_COLUMNS)}
        raise SystemExit(f"{path}: not a DRA tax rate or tax rate calculation workbook.")
    text = content.decode("utf-8-sig", errors="replace")
    if "Cohort Graduation" in text[:500]:
        year, rows = graduation(text)
        return "graduation", {year: rows}
    if text.startswith(("DenominatorType,", "yearid,")) and "Above prof%" in text[:500]:
        return "assessment", assessment(text)
    if re.search(r"COST PER PUPIL BY DISTRICT", text[:3000], re.I):
        year, rows = cost_per_pupil(text)
        return "cost_per_pupil", {year: rows}
    if text.startswith("SUMLEV,"):
        return "population", population(text)
    raise SystemExit(f"{path}: not a file this command knows. See the list at the top of pipeline/states/nh/extract.py.")


def merge(saved: dict, kind: str, years: dict[int, dict]) -> None:
    """Add each year's rows. The tax files come in pairs (rates, calculation) that fill in the same records."""
    for year, rows in years.items():
        key = str(year)
        if kind == "tax":
            existing = saved["years"].setdefault(key, {})
            for name, record in rows.items():
                existing[name] = {**existing.get(name, {}), **record}
            saved["years"][key] = dict(sorted(existing.items()))
        else:
            saved["years"][key] = dict(sorted(rows.items()))


def fetch_population(now: datetime) -> str:
    """The newest Census city and town estimates for New Hampshire (census.gov allows scripts)."""
    import requests
    for year in range(now.year, now.year - 3, -1):
        response = requests.get(POPULATION_URL.format(year=year), timeout=60,
                                headers={"User-Agent": "publick.org (+https://publick.org)"})
        if response.ok and response.text.startswith("SUMLEV,"):
            print(f"Census estimates, vintage {year}")
            return response.content.decode("latin-1")
    raise SystemExit("No Census city and town estimates found for the last three years.")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("files", nargs="*", type=Path, help="downloaded DRA and Department of Education files")
    parser.add_argument("--population", action="store_true", help="also fetch the Census population estimates")
    args = parser.parse_args()
    if not args.files and not args.population:
        parser.error("name the downloaded files, or --population")
    now = datetime.now(timezone.utc)
    found: dict[str, list] = {}
    for path in args.files:
        kind, years = read(path)
        found.setdefault(kind, []).append(years)
        print(f"{path.name}: {kind}, {', '.join(str(y) for y in sorted(years))}")
    if args.population:
        found.setdefault("population", []).append(population(fetch_population(now)))
    for kind, batches in found.items():
        saved = load(kind)
        for years in batches:
            merge(saved, kind, years)
        save(kind, saved, now)
        print(f"Saved figures/{kind}.json: {', '.join(saved['years'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
