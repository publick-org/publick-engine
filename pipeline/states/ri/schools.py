"""Rhode Island school district figures for the schools page, from the Rhode
Island Department of Education (RIDE). Each as published, for the district and
the state:

- the four-year cohort graduation rate (report card data files,
  GraduationRates_<year>.xlsx);
- spending per pupil, all sources (report card data files,
  Finance_<year>.xlsx: RIDE's "Total Per Pupil");
- the share of students in grades 3 to 8 meeting or exceeding expectations on
  RICAS, the Rhode Island Comprehensive Assessment System, in English and math
  (RIDE's assessment data portal, ADP, whose export is a tab-separated table);
- the share of students chronically absent, for the district only: the
  all-students figure is in the accountability data file
  (Accountability_<year>.xlsx, "LEA Indicator Data"), which has no statewide
  row, and RIDE's chronic absenteeism file has statewide figures only for groups
  of students, not all of them. So the state's figure is None.

Every report card data file is listed on RIDE's Data Files page (DATA_FILES),
in a folder per report card year ("202425"), which holds the previous school
year's graduation rate and spending (the class of 2024, 2023-24) and the
report card year's own absenteeism. Layouts drift between years, so columns are
found by their headers, and each figure's year is read from the file where it
says. Each run reads only the files of report card years not read yet, and the
newest year's again (RIDE replaces files with corrected versions), and keeps
the years already saved, up to YEARS_KEPT. The ADP export covers every year
RICAS has been given, so it's asked for in one request per subject.

Years are school years' ending years (2024-25 is 2025), a class's graduation
year, and a test's spring. Writes data/schools/schools.json. Run by
python -m pipeline.fetch_schools.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import urlencode
from zoneinfo import ZoneInfo

import openpyxl

from pipeline.fetch_meetings import save_json
from pipeline.http import FetchError, PoliteClient
from pipeline.i18n import N_, _
from pipeline.rhythms import Part, Rhythm, latest_year, on

REPORT_CARD = "https://reportcard.ride.ri.gov"
DATA_FILES = f"{REPORT_CARD}/DataFiles"
ADP = "https://www3.ride.ri.gov/ADP"
ADP_YEARS = f"{ADP}/Default/GetYearsByAssessment"
ADP_EXPORT = f"{ADP}/Default/Export"
# ADP's assessments: RICAS English language arts and math.
ASSESSMENTS = {"ricas_ela": 5, "ricas_math": 6}
# ADP's code for the state, as a district and as what a district is compared with.
STATE_LEA = "00"
YEARS_KEPT = 8
# Each report card data file's name, and the measure it holds.
REPORT_FILES = {"graduation": "GraduationRates", "spending": "Finance", "absenteeism": "Accountability"}

SOURCE_PAGES = {"graduation": DATA_FILES, "spending": DATA_FILES, "absenteeism": DATA_FILES,
                "ricas_ela": f"{ADP}/", "ricas_math": f"{ADP}/"}

# RIDE's 2024-25 report card (with the class of 2024's graduation rate, 2023-24
# spending, and 2024-25 absenteeism) was posted in November 2025; spring 2025's
# RICAS results were in the ADP by fall 2025.
RHYTHM = Rhythm(N_("School figures"), "schools/schools.json", "Fetch school figures", "yearly", (
    Part(latest_year("measures.graduation.years", "year"), on(12, years_after=1),
         lambda y: _("Class of {year} graduation rate").format(year=y)),
    Part(latest_year("measures.ricas_ela.years", "year"), on(10), lambda y: _("Spring {year} state test results").format(year=y)),
    Part(latest_year("measures.absenteeism.years", "year"), on(12),
         lambda y: _("{years} attendance").format(years=f"{y - 1}–{y % 100:02d}")),
    Part(latest_year("measures.spending.years", "year"), on(12, years_after=1),
         lambda y: _("{years} spending per pupil").format(years=f"{y - 1}–{y % 100:02d}")),
))


def client(config: dict) -> PoliteClient:
    return PoliteClient(config["site"]["user_agent"], delay=1.0, timeout=120)


def code(value) -> str:
    """A district code as RIDE's files write it, '01' to '99': config may say 32 or "32"."""
    return str(value).strip().zfill(2)


def ending_year(label) -> int | None:
    """'2023-24' or '2016-2017' -> the school year's ending year."""
    m = re.fullmatch(r"(\d{4})-(\d{2}|\d{4})", str(label or "").strip())
    return int(m.group(1)) + 1 if m else None


def norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def number(value) -> float | None:
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError:
        return None


# ---- The report card data files ----

def report_files(html: str) -> dict[str, dict[str, str]]:
    """{measure: {report card year ("202425"): file's address}} from the Data Files page."""
    found: dict[str, dict[str, str]] = {}
    for measure, name in REPORT_FILES.items():
        for m in re.finditer(rf'href="(/(\d{{6}})/datafiles/{name}_(\d{{6}})[^"/]*\.xlsx)"', html, re.I):
            if m.group(2) == m.group(3):
                found.setdefault(measure, {})[m.group(2)] = REPORT_CARD + m.group(1)
    return found


def sheet_rows(book, *names: str, header_rows: int = 1) -> tuple[list[list[str]], list[tuple]]:
    """The first sheet with one of these names (any case): its header rows, normalized, and the rows after."""
    sheet = next((book[s] for s in book.sheetnames if s.lower() in names), None)
    if sheet is None:
        return [], []
    rows = list(sheet.iter_rows(values_only=True))
    return [[norm(c) for c in row] for row in rows[:header_rows]], rows[header_rows:]


def column(headers: list[list[str]], *names: str) -> int | None:
    """The index of the column one of whose header cells is one of these names."""
    for header in headers:
        for name in names:
            if name in header:
                return header.index(name)
    return None


def workbook(content: bytes):
    return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)


def graduation(content: bytes, district: str) -> tuple[dict[int, float], dict[int, float]]:
    """({class year: district's four-year rate}, {class year: state's}) from a GraduationRates file."""
    book = workbook(content)
    out = []
    for names, is_state in ((("lea", "leas"), False), (("state",), True)):
        headers, rows = sheet_rows(book, *names)
        year, rate, kind, group = (column(headers, *n) for n in (("schoolyear", "schoolyearvalue"),
                                                                    ("cohortgraduationrate4",), ("graduationrate",),
                                                                    ("groupcode",)))
        lea, grouping = column(headers, "leacode"), column(headers, "stategrouping")
        found = {}
        if None in (year, rate, kind, group) or (lea is None and not is_state):
            rows = []
        for row in rows:
            if (norm(row[kind]) != "4 year rates" or norm(row[group]) != "all"
                    or (is_state and grouping is not None and norm(row[grouping]) != "rhode island")
                    or (not is_state and (row[lea] is None or code(row[lea]) != district))):
                continue
            if (y := ending_year(row[year])) and (value := number(row[rate])) is not None:
                found[y] = value
        out.append(found)
    return out[0], out[1]


def spending(content: bytes, district: str) -> tuple[dict[int, float], dict[int, float]]:
    """({school year: district's total per pupil}, {school year: state's}) from a Finance file."""
    book = workbook(content)
    out = []
    for names, is_state in ((("lea", "leas"), False), (("state",), True)):
        headers, rows = sheet_rows(book, *names)
        year, value, kind, lea = (column(headers, n) for n in ("schoolyear", "perpupil", "rowtype", "leacode"))
        found = {}
        if None in (year, value, kind) or (lea is None and not is_state):
            rows = []
        for row in rows:
            if norm(row[kind]) != "total per pupil":
                continue
            if not is_state and (row[lea] is None or code(row[lea]) != district):
                continue
            if (y := ending_year(row[year])) and (n := number(row[value])) is not None:
                found[y] = round(n, 2)
        out.append(found)
    return out[0], out[1]


def absenteeism(content: bytes, district: str) -> float | None:
    """The district's all-students chronic absenteeism rate from an Accountability file, or None when
    the file has none (the years before 2021-22 have no district sheet). The sheet has two header rows,
    RIDE's labels and its column names, in either order."""
    headers, rows = sheet_rows(workbook(content), "lea indicator data", header_rows=2)
    rate = column(headers, "% student chronic abs", "student_ca_pct")
    lea = column(headers, "district code", "distcode")
    group, group_code = column(headers, "group"), column(headers, "group code", "groupcode")
    if rate is None or lea is None:
        return None
    for row in rows:
        if row[lea] is None or code(row[lea]) != district:
            continue
        if (group is not None and norm(row[group]) == "all students") or \
                (group is None and group_code is not None and norm(row[group_code]) == "all"):
            return number(row[rate])
    return None


# ---- RICAS, from the assessment data portal ----

def adp_years(client, assessment: int) -> list[str]:
    response = client.get(f"{ADP_YEARS}?{urlencode({'assessment': assessment})}")
    return [y["code"] for y in json.loads(response.text)]


def adp_url(assessment: int, district: str, years: list[str]) -> str:
    return ADP_EXPORT + "?" + urlencode({
        "assessment": assessment, "lea": district, "sch": 0, "grade": 0, "supergroup": 0,
        "schYear": ",".join(years), "compareWith": STATE_LEA, "ADA": "false", "Growth": "false",
        "Performance": "false"}, safe=",")


def parse_adp(text: str) -> tuple[dict[int, float], dict[int, float]]:
    """({spring: district's share meeting or exceeding expectations}, {spring: state's}) from an ADP
    export: the district's rows, then the state's ("Statewide"), for all schools, grades, and students."""
    if "Percent_Meeting_or_Exceeding_Expectations" not in text[:1000]:
        raise FetchError("the assessment portal's export has no results table: " + text[:200])
    town, state = {}, {}
    for r in csv.DictReader(io.StringIO(text), delimiter="\t"):
        if (r.get("School"), r.get("Grade"), r.get("Group")) != ("All Schools", "All Grades", "All Groups"):
            continue
        year, value = ending_year(r["School_Year"]), number(r["Percent_Meeting_or_Exceeding_Expectations"])
        if year and value is not None:
            (state if r["District"] == "Statewide" else town)[year] = value
    return town, state


# ---- The source ----

def merged(saved: list[dict], town: dict[int, float], state: dict[int, float | None]) -> list[dict]:
    """The saved years with the newly read ones, which replace a saved year's figures, oldest first."""
    years = {y["year"]: y for y in saved}
    for year, value in town.items():
        years[year] = {"year": year, "town": value, "state": state.get(year)}
    return [years[y] for y in sorted(years)][-YEARS_KEPT:]


def run(config: dict, client, data_dir: Path, now: datetime | None = None, force: bool = False) -> dict:
    now = now or datetime.now(ZoneInfo(config["site"]["timezone"]))
    schools = config["schools"]
    district = code(schools["ride_district"])
    path = data_dir / "schools" / "schools.json"
    saved = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    saved_measures, read_before = saved.get("measures", {}), saved.get("report_cards", {})

    def years(name: str) -> list[dict]:
        return saved_measures.get(name, {}).get("years", [])

    files = report_files(client.get(DATA_FILES).text)
    if not files.get("graduation"):
        raise FetchError("no graduation rate files on RIDE's Data Files page")
    measures, report_cards = {}, {}
    for measure, by_year in files.items():
        # Report card years old enough that their figures have dropped off aren't read.
        recent = [y for y in sorted(by_year) if int(y[:4]) + 1 >= now.year - YEARS_KEPT]
        newest = recent[-1] if recent else None
        wanted = [y for y in recent if y not in read_before.get(measure, []) or y == newest or force]
        town, state = {}, {}
        for year in wanted:
            content = client.get(by_year[year]).content
            if measure == "absenteeism":
                if (value := absenteeism(content, district)) is not None:
                    town[int(year[:4]) + 1] = value
            else:
                t, s = (graduation if measure == "graduation" else spending)(content, district)
                town.update(t)
                state.update(s)
        report_cards[measure] = sorted(set(read_before.get(measure, [])) | set(wanted))
        measures[measure] = {"source_url": SOURCE_PAGES[measure], "years": merged(years(measure), town, state)}
    # A kind of file missing from the page this run (renamed, or briefly taken down) keeps what's saved.
    for measure in SOURCE_PAGES:
        if measure not in measures and measure in saved_measures:
            measures[measure], report_cards[measure] = saved_measures[measure], read_before.get(measure, [])

    for measure, assessment in ASSESSMENTS.items():
        town, state = parse_adp(client.get(adp_url(assessment, district, adp_years(client, assessment))).text)
        if not town:
            raise FetchError(f"no RICAS results for district {district}")
        measures[measure] = {"source_url": SOURCE_PAGES[measure], "years": merged(years(measure), town, state)}

    for measure in ("graduation", "spending"):
        if not measures.get(measure, {}).get("years"):
            raise FetchError(f"no {measure} figures for district {district} in RIDE's files")
    result = {
        "updated_at": now.isoformat(timespec="seconds"),
        "source": "Rhode Island Department of Education",
        "district": schools["district_name"],
        "report_cards": report_cards,
        "measures": measures,
    }
    save_json(path, result)
    return {name: m["years"][-1] if m["years"] else None for name, m in measures.items()}
